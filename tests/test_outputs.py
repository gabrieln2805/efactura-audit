import io
import os

import pytest
from openpyxl import load_workbook

from efactura_audit.cli import main
from efactura_audit.report import excel_report, pdf_report


@pytest.mark.parametrize("lang,sheet", [("ro", "Semnalări"), ("en", "Flags")])
def test_excel_report(demo, lang, sheet):
    wb = load_workbook(io.BytesIO(excel_report(demo["store"], lang)))
    assert sheet in wb.sheetnames
    ws = wb[sheet]
    assert ws.max_row - 1 == len(demo["store"].flags())


def test_pdf_report(demo):
    data = pdf_report(demo["store"], "ro")
    assert data[:4] == b"%PDF" and len(data) > 20_000


def test_cli_demo_status_report(tmp_path, capsys):
    ws = str(tmp_path / "ws")
    main(["-w", ws, "demo"])
    out = capsys.readouterr().out
    assert "Recall:" in out and "invoices audited" in out
    main(["-w", ws, "report", "--lang", "en", "--format", "xlsx"])
    assert any(p.suffix == ".xlsx" for p in (tmp_path / "ws" / "reports").iterdir())
    main(["-w", ws, "status"])
    assert '"invoices"' in capsys.readouterr().out
    with pytest.raises(SystemExit):
        main(["-w", ws, "ask", "Câte facturi?", "--model", "not-tested"])


def test_streamlit_app_smoke(demo):
    AppTest = pytest.importorskip("streamlit.testing.v1").AppTest
    os.environ["EFACTURA_HOME"] = str(demo["root"])
    app = os.path.join(os.path.dirname(__file__), "..", "src", "efactura_audit", "app", "streamlit_app.py")
    at = AppTest.from_file(app, default_timeout=90)
    at.run()
    assert not at.exception
    assert [m.value for m in at.metric][0] == f"{demo['summary']['invoices']}"
    at.sidebar.segmented_control[0].set_value("en").run()
    assert not at.exception and at.metric[0].label == "Invoices"
    # assistant with an untested local model while Ollama is not running: a clear error, no crash
    at.toggle[1].set_value(True).run()
    at.chat_input[0].set_value("How many invoices?").run()
    assert not at.exception


def test_free_port_skips_busy_port():
    import socket
    from efactura_audit.cli import free_port
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        s.listen()
        busy = s.getsockname()[1]
        assert free_port(busy) != busy

"""Controller reports: Excel workbook and PDF, in Romanian or English.

Flags are presented as prompts to check, never as conclusions — the disclaimer is on every report.
"""
from __future__ import annotations

import io
import json
from datetime import datetime
from pathlib import Path

import pandas as pd

from .checks.base import CHECK_LABELS, lei
from .store import Store

SERIES_BLUE = "#2a78d6"
SEVERITY_COLORS = {"high": "#d03b3b", "medium": "#ec835a", "low": "#fab219"}

T = {
    "en": {
        "title": "Clarifact · Supplier invoice checks",
        "period": "Period", "invoices": "Invoices", "suppliers": "Suppliers", "spend": "Supplier spend (incl. VAT)",
        "flags": "Flags", "at_risk": "Amount to review", "by_check": "Flags by check", "by_sev": "Flags by severity",
        "top": "Flags to review, most important first", "generated": "Generated",
        "severity": "Severity", "check": "Check", "supplier": "Supplier", "invoices_col": "Invoice(s)",
        "date": "Date", "amount": "Amount to review (lei)", "finding": "Finding", "todo": "What to check",
        "status": "Status", "rule": "Rule", "evidence": "Evidence", "flag_id": "Flag ID",
        "sev": {"high": "High", "medium": "Medium", "low": "Low"},
        "disclaimer": ("Flags are prompts to check, not conclusions. They are generated automatically from the "
                       "e-invoices and contain no tax or accounting advice. Confirm each one against contracts, "
                       "delivery notes and the supplier before acting."),
        "summary_sheet": "Summary", "flags_sheet": "Flags", "suppliers_sheet": "Suppliers",
        "sup_cols": ["Supplier", "CUI", "Invoices", "Spend (lei)", "Flags", "Amount to review (lei)"],
        "company": "Company", "more": "{n} more flags in the Excel report.",
    },
    "ro": {
        "title": "Clarifact · Verificarea facturilor de la furnizori",
        "period": "Perioada", "invoices": "Facturi", "suppliers": "Furnizori", "spend": "Achiziții (cu TVA)",
        "flags": "Semnalări", "at_risk": "Sumă de verificat", "by_check": "Semnalări pe tip de verificare",
        "by_sev": "Semnalări pe grad de importanță", "top": "Semnalări de verificat, în ordinea importanței",
        "generated": "Generat", "severity": "Importanță", "check": "Verificare", "supplier": "Furnizor",
        "invoices_col": "Factura/facturile", "date": "Data", "amount": "Sumă de verificat (lei)",
        "finding": "Constatare", "todo": "Ce trebuie verificat", "status": "Stare", "rule": "Regulă",
        "evidence": "Dovezi", "flag_id": "ID semnalare",
        "sev": {"high": "Ridicată", "medium": "Medie", "low": "Scăzută"},
        "disclaimer": ("Semnalările sunt indicii de verificat, nu concluzii. Sunt generate automat din facturile "
                       "electronice și nu reprezintă consultanță fiscală sau contabilă. Confirmați fiecare semnalare "
                       "cu contractele, avizele de însoțire și furnizorul înainte de a acționa."),
        "summary_sheet": "Rezumat", "flags_sheet": "Semnalări", "suppliers_sheet": "Furnizori",
        "sup_cols": ["Furnizor", "CUI", "Facturi", "Achiziții (lei)", "Semnalări", "Sumă de verificat (lei)"],
        "company": "Companie", "more": "Încă {n} semnalări în raportul Excel.",
    },
}

STATUS = {"en": {"open": "Open", "confirmed": "Confirmed", "dismissed": "Dismissed"},
          "ro": {"open": "Deschisă", "confirmed": "Confirmată", "dismissed": "Respinsă"}}


def _check_label(check: str, lang: str) -> str:
    en, ro = CHECK_LABELS.get(check, (check, check))
    return ro if lang == "ro" else en


def gather(store: Store) -> dict:
    inv = store.df("SELECT * FROM fct_invoices")
    flags = store.flags(include_dismissed=False)
    sup = store.df(
        """SELECT supplier_name, supplier_cui, invoices AS n, total_spend_ron AS spend, flags AS nflags,
                  amount_to_review_ron AS risk
           FROM dim_suppliers ORDER BY spend DESC"""
    )
    buyer = inv["buyer_name"].mode().iloc[0] if not inv.empty else ""
    return {
        "invoices": inv,
        "flags": flags,
        "suppliers": sup,
        "buyer": buyer,
        "period": (str(inv["issue_date"].min())[:10], str(inv["issue_date"].max())[:10]) if not inv.empty else ("", ""),
        "spend": float(inv["spend_ron"].sum()) if not inv.empty else 0.0,
    }


def _todo(evidence: str | dict, lang: str) -> str:
    ev = json.loads(evidence) if isinstance(evidence, str) else evidence
    notes = ev.get(f"notes_{lang}") or []
    base = ev.get(f"what_to_check_{lang}", "")
    return " ".join([*notes, base]).strip()


def _evidence_text(evidence: str | dict) -> str:
    ev = json.loads(evidence) if isinstance(evidence, str) else evidence
    ev = {k: v for k, v in ev.items() if not k.startswith(("what_to_check", "notes_")) and not k.endswith("_full")}
    return json.dumps(ev, ensure_ascii=False, default=str)[:3000]


def flags_table(flags: pd.DataFrame, lang: str) -> pd.DataFrame:
    t = T[lang]
    if flags.empty:
        return pd.DataFrame(columns=[t["severity"], t["check"], t["supplier"], t["invoices_col"], t["date"],
                                     t["amount"], t["finding"], t["todo"], t["status"], t["rule"], t["flag_id"]])
    return pd.DataFrame({
        "#": range(1, len(flags) + 1),
        t["severity"]: flags["severity"].map(t["sev"]),
        t["check"]: flags["check_name"].map(lambda c: _check_label(c, lang)),
        t["supplier"]: flags["supplier_name"],
        t["invoices_col"]: flags["invoice_numbers"].map(lambda v: ", ".join(json.loads(v))[:120]),
        t["date"]: flags["issue_date"].astype(str).str[:10],
        t["amount"]: flags["amount_at_risk_ron"].round(2),
        t["finding"]: flags[f"title_{lang}"],
        t["todo"]: flags["evidence"].map(lambda e: _todo(e, lang)),
        t["status"]: flags["status"].map(STATUS[lang]),
        t["rule"]: flags["check_name"] + "/" + flags["rule"],
        t["evidence"]: flags["evidence"].map(_evidence_text),
        t["flag_id"]: flags["flag_id"],
    })


# --------------------------------------------------------------------------------------- Excel
def excel_report(store: Store, lang: str = "ro") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    t = T[lang]
    data = gather(store)
    flags = data["flags"]
    wb = Workbook()
    ws = wb.active
    ws.title = t["summary_sheet"]
    bold = Font(bold=True)
    ws["A1"] = t["title"]
    ws["A1"].font = Font(bold=True, size=16)
    rows = [
        (t["company"], data["buyer"]),
        (t["period"], f"{data['period'][0]} – {data['period'][1]}"),
        (t["invoices"], len(data["invoices"])),
        (t["suppliers"], len(data["suppliers"])),
        (t["spend"], round(data["spend"], 2)),
        (t["flags"], len(flags)),
        (t["at_risk"], round(float(flags["amount_at_risk_ron"].fillna(0).sum()), 2) if not flags.empty else 0),
        (t["generated"], datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = bold
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, float):
            c.number_format = "#,##0.00"
    r = len(rows) + 5
    ws.cell(row=r, column=1, value=t["by_check"]).font = bold
    for check in CHECK_LABELS:
        r += 1
        n = int((flags["check_name"] == check).sum()) if not flags.empty else 0
        ws.cell(row=r, column=1, value=_check_label(check, lang))
        ws.cell(row=r, column=2, value=n)
    r += 2
    ws.cell(row=r, column=1, value=t["by_sev"]).font = bold
    for sev in ("high", "medium", "low"):
        r += 1
        n = int((flags["severity"] == sev).sum()) if not flags.empty else 0
        ws.cell(row=r, column=1, value=t["sev"][sev])
        ws.cell(row=r, column=2, value=n)
    r += 2
    c = ws.cell(row=r, column=1, value=t["disclaimer"])
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 3, end_column=4)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20

    # Flags sheet
    ft = flags_table(flags, lang)
    ws2 = wb.create_sheet(t["flags_sheet"])
    ws2.append(list(ft.columns))
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="104281")
    fills = {t["sev"][k]: PatternFill("solid", fgColor=v.lstrip("#")) for k, v in SEVERITY_COLORS.items()}
    for row in ft.itertuples(index=False):
        ws2.append([None if (isinstance(x, float) and pd.isna(x)) else x for x in row])
        sev_cell = ws2.cell(row=ws2.max_row, column=2)
        if sev_cell.value in fills:
            sev_cell.fill = fills[sev_cell.value]
            sev_cell.font = Font(bold=True, color="FFFFFF" if sev_cell.value != t["sev"]["low"] else "000000")
    widths = [5, 11, 18, 30, 24, 11, 16, 60, 60, 12, 26, 60, 16]
    for i, w in enumerate(widths[: len(ft.columns)], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[6].number_format = "#,##0.00"
    ws2.freeze_panes = "B2"
    ws2.auto_filter.ref = ws2.dimensions

    # Suppliers sheet
    ws3 = wb.create_sheet(t["suppliers_sheet"])
    ws3.append(t["sup_cols"])
    for cell in ws3[1]:
        cell.font = bold
    for s in data["suppliers"].itertuples(index=False):
        ws3.append([s.supplier_name, s.supplier_cui, int(s.n), round(float(s.spend or 0), 2), int(s.nflags), round(float(s.risk), 2)])
    for col, w in zip("ABCDEF", (40, 14, 10, 18, 10, 22)):
        ws3.column_dimensions[col].width = w
    for row in ws3.iter_rows(min_row=2):
        row[3].number_format = "#,##0.00"
        row[5].number_format = "#,##0.00"
    ws3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------------------- PDF
def _register_font() -> tuple[str, str]:
    """A TTF with Romanian diacritics (ș ț ă î â). matplotlib ships DejaVu Sans."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = []
    try:
        import matplotlib
        base = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        candidates.append((base / "DejaVuSans.ttf", base / "DejaVuSans-Bold.ttf"))
    except Exception:  # noqa: BLE001
        pass
    candidates += [
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    ]
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            if "AuditSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("AuditSans", str(regular)))
                pdfmetrics.registerFont(TTFont("AuditSans-Bold", str(bold)))
            return "AuditSans", "AuditSans-Bold"
    return "Helvetica", "Helvetica-Bold"


def _bar_png(counts: list[tuple[str, int]], title: str) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    labels = [c[0] for c in counts][::-1]
    values = [c[1] for c in counts][::-1]
    fig, ax = plt.subplots(figsize=(6.4, 2.2), dpi=200)
    bars = ax.barh(labels, values, color=SERIES_BLUE, height=0.55)
    for b, v in zip(bars, values):
        ax.text(b.get_width() + max(values + [1]) * 0.01, b.get_y() + b.get_height() / 2, str(v),
                va="center", fontsize=8, color="#0b0b0b")
    ax.set_title(title, fontsize=9, loc="left", color="#0b0b0b")
    for s in ("top", "right", "bottom"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#c3c2b7")
    ax.tick_params(axis="y", labelsize=8, colors="#52514e", length=0)
    ax.set_xticks([])
    ax.set_xlim(0, max(values + [1]) * 1.15)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def pdf_report(store: Store, lang: str = "ro", max_flags: int = 60) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    t = T[lang]
    regular, bold = _register_font()
    data = gather(store)
    flags = data["flags"]
    styles = {
        "h1": ParagraphStyle("h1", fontName=bold, fontSize=18, leading=22, spaceAfter=4),
        "h2": ParagraphStyle("h2", fontName=bold, fontSize=12, leading=15, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("b", fontName=regular, fontSize=9, leading=12),
        "small": ParagraphStyle("s", fontName=regular, fontSize=7.5, leading=9.5, alignment=TA_LEFT),
        "smallb": ParagraphStyle("sb", fontName=bold, fontSize=7.5, leading=9.5, textColor=colors.white),
        "muted": ParagraphStyle("m", fontName=regular, fontSize=8, leading=10, textColor=colors.HexColor("#52514e")),
    }
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=t["title"])
    story = [Paragraph(t["title"], styles["h1"]),
             Paragraph(f"{data['buyer']} · {t['period']}: {data['period'][0]} – {data['period'][1]} · "
                       f"{t['generated']}: {datetime.now():%Y-%m-%d %H:%M}", styles["muted"]),
             Spacer(1, 6)]

    at_risk = float(flags["amount_at_risk_ron"].fillna(0).sum()) if not flags.empty else 0.0
    kpis = [[t["invoices"], t["suppliers"], t["spend"], t["flags"], t["at_risk"]],
            [f"{len(data['invoices']):,}", f"{len(data['suppliers'])}", lei(data["spend"], lang),
             f"{len(flags)}", lei(at_risk, lang)]]
    kt = Table(kpis, colWidths=[52 * mm] * 5)
    kt.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), regular), ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#52514e")),
        ("FONTNAME", (0, 1), (-1, 1), bold), ("FONTSIZE", (0, 1), (-1, 1), 14),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 8), ("LINEBELOW", (0, 1), (-1, 1), 0.5, colors.HexColor("#c3c2b7")),
    ]))
    story += [kt, Spacer(1, 6)]

    counts = [(_check_label(c, lang), int((flags["check_name"] == c).sum()) if not flags.empty else 0) for c in CHECK_LABELS]
    sev_counts = {s: int((flags["severity"] == s).sum()) if not flags.empty else 0 for s in ("high", "medium", "low")}
    chart = Image(io.BytesIO(_bar_png(counts, t["by_check"])), width=150 * mm, height=51 * mm)
    sev_rows = [[Paragraph(f"<b>{t['by_sev']}</b>", styles["body"]), ""]]
    for s in ("high", "medium", "low"):
        sev_rows.append([Paragraph(f"<font color='{SEVERITY_COLORS[s]}'>■</font> {t['sev'][s]}", styles["body"]), str(sev_counts[s])])
    st = Table(sev_rows, colWidths=[50 * mm, 20 * mm])
    st.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), regular), ("FONTSIZE", (0, 0), (-1, -1), 9),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    row = Table([[chart, st]], colWidths=[160 * mm, 80 * mm])
    row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [row, Paragraph(t["disclaimer"], styles["muted"]), Spacer(1, 4)]

    story.append(Paragraph(t["top"], styles["h2"]))
    header = ["#", t["severity"], t["check"], t["supplier"], t["invoices_col"], t["amount"], t["finding"], t["todo"]]
    rows = [[Paragraph(f"<b>{h}</b>", styles["smallb"]) for h in header]]
    shown = flags.head(max_flags)
    for i, f in enumerate(shown.itertuples(index=False), 1):
        nums = ", ".join(json.loads(f.invoice_numbers))[:60]
        amt = "" if pd.isna(f.amount_at_risk_ron) else f"{f.amount_at_risk_ron:,.2f}"
        if lang == "ro" and amt:
            amt = amt.replace(",", "_").replace(".", ",").replace("_", ".")
        rows.append([
            Paragraph(str(i), styles["small"]),
            Paragraph(f"<font color='{SEVERITY_COLORS[f.severity]}'>■</font> {t['sev'][f.severity]}", styles["small"]),
            Paragraph(_check_label(f.check_name, lang), styles["small"]),
            Paragraph(str(f.supplier_name or ""), styles["small"]),
            Paragraph(nums, styles["small"]),
            Paragraph(amt, styles["small"]),
            Paragraph(getattr(f, f"title_{lang}"), styles["small"]),
            Paragraph(_todo(f.evidence, lang), styles["small"]),
        ])
    widths = [9, 21, 22, 36, 30, 22, 60, 69]
    tbl = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#104281")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f3")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e3e2dc")),
        ("ALIGN", (5, 1), (5, -1), "RIGHT"),
    ]))
    story.append(tbl)
    if len(flags) > max_flags:
        story += [Spacer(1, 4), Paragraph(t["more"].format(n=len(flags) - max_flags), styles["muted"])]
    doc.build(story)
    return buf.getvalue()


def write_reports(store: Store, out_dir: Path, lang: str = "ro", formats: tuple[str, ...] = ("xlsx", "pdf")) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    paths = []
    if "xlsx" in formats:
        p = out_dir / f"clarifact-{stamp}-{lang}.xlsx"
        p.write_bytes(excel_report(store, lang))
        paths.append(p)
    if "pdf" in formats:
        p = out_dir / f"clarifact-{stamp}-{lang}.pdf"
        p.write_bytes(pdf_report(store, lang))
        paths.append(p)
    return paths
